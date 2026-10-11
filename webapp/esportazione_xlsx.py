"""Generazione di file XLSX per l'export di storico e anagrafica clienti."""
import io

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

_COLORE_INTESTAZIONE = "0F6B4C"  # stesso verde dell'accento usato nel sito


def crea_foglio_xlsx(titolo_foglio: str, intestazioni: list, righe: list,
                      larghezze: list = None, formati_colonna: dict = None) -> io.BytesIO:
    """
    Crea un workbook XLSX con un solo foglio.

    - `righe`: lista di liste di valori (stringhe, numeri o oggetti date),
      nello stesso ordine delle `intestazioni`.
    - `larghezze`: larghezza (in "caratteri") di ciascuna colonna, opzionale.
    - `formati_colonna`: {indice_colonna (1-based): formato numerico}, es.
      {4: '#,##0.00', 2: 'DD/MM/YYYY'}, applicato a tutte le righe dati.
    """
    cartella = Workbook()
    foglio = cartella.active
    foglio.title = titolo_foglio

    foglio.append(intestazioni)
    for cella in foglio[1]:
        cella.font = Font(bold=True, color="FFFFFF")
        cella.fill = PatternFill("solid", fgColor=_COLORE_INTESTAZIONE)

    for riga in righe:
        foglio.append(riga)

    for indice_colonna, formato in (formati_colonna or {}).items():
        lettera = get_column_letter(indice_colonna)
        for cella in foglio[lettera][1:]:
            cella.number_format = formato

    for indice, valore in enumerate(larghezze or [], start=1):
        foglio.column_dimensions[get_column_letter(indice)].width = valore

    foglio.freeze_panes = "A2"

    buffer = io.BytesIO()
    cartella.save(buffer)
    buffer.seek(0)
    return buffer
